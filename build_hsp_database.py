#!/usr/bin/env python3
"""Build comprehensive HSP database from multiple public data sources.

Data sources (in priority order):
1. Original curated data - 99 solvents + 30 polymers from Hansen's Handbook (2007)
2. Mendeley Data (Langner & Brabec 2022) - 499 solvents, CC BY 4.0
3. SolvPred (Fang et al.) - 249 solvents with physical properties
4. Accudyne Test - 89 solvents + 117 polymers from Hansen's Handbook
5. Wolfram Data Repository (Schrier 2020) - 211 solvents
6. Pang et al. 2024 - 1,183 compounds from HSPiP database
7. HSPiP Software Database - 1,218 solvents with CAS numbers
8. Hansen Appendix A.2 - 458 polymers/materials from Hansen's Handbook appendix

Usage:
    python build_hsp_database.py
"""

import csv
import json
import os
import re
import sys

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RAW_DIR = os.path.join(BASE_DIR, "data", "raw")
OUT_DIR = os.path.join(BASE_DIR, "data", "processed")

SOURCE_URLS = {
    "handbook": "https://hansen-solubility.com (Hansen Handbook 2007)",
    "hspip": "https://www.hansen-solubility.com/HSPiP/ (HSPiP Software Database)",
    "hansen_a2": "https://hansen-solubility.com (Hansen Handbook Appendix Table A.2)",
    "mendeley": "https://data.mendeley.com/datasets/b4dmjzk8w6/1",
    "solvpred": "https://github.com/xueannafang/hsp_toolkit_solv_pred_v_2.0",
    "accudyne": "https://www.accudynetest.com/solubility_table.html",
    "wolfram": "https://datarepository.wolframcloud.com/resources/JoshuaSchrier_Hansen-Solubility-Parameters/",
    "pang2024": "https://github.com/jiayunpang/hsp_embedding",
}


def normalize_name(name):
    """Normalize chemical name for dedup comparison."""
    if not name:
        return ""
    n = name.strip().lower()
    n = re.sub(r"\s+", " ", n)
    # Remove trailing parenthetical descriptions
    n = re.sub(r"\s*\([^)]*\)\s*$", "", n)
    # Strip punctuation and whitespace for comparison
    n = n.replace("-", "").replace(",", "").replace(" ", "").replace(".", "")
    return n


# --- Hansen A1 Name Splitting ---
# Legacy helper: extract_common_name() was used when the old HSP_A1_Final.csv
# had "CommonName IUPACName" concatenated. Now table_a1.csv has separate columns,
# but these functions are kept for any other callers.

# IUPAC carbon-chain stems
_IUPAC_STEMS = re.compile(
    r"(meth|eth(?!yl)|prop|but(?!yl)|pent|hex|hept|oct|non(?!yl)|dec|"
    r"undec|dodec|benz[eo]|naphthal|oxir|thiir|azirid|furan|pyrrol|"
    r"pyrid|thiophen|imidazol|morpholin|quinol|chromen|oxepan|thiet)",
    re.I,
)

# Systematic suffixes (without preceding hyphen)
_IUPAC_SUFFIX = re.compile(
    r"(ane|ene|yne|ol|al|one|amine|amide|nitrile|thiol|diol|dione|"
    r"triol|anone|enol|anide|oate|oic|ide|ate|ole)$",
    re.I,
)


def _looks_like_iupac_word(word):
    """Check if a word is a systematic IUPAC-style name (with hyphens/digits)."""
    # Contains hyphen-digit pattern (Pent-4-enoic, But-2-enal)
    if re.search(r"-\d", word):
        return True
    # Hyphenated systematic name (Bromo-benzene, Chloro-acetaldehyde)
    if re.match(r"^[A-Z][a-z]+-[a-z]", word):
        return True
    # Di/Tri/etc prefix with hyphens (Dichloro-fluoro-methane)
    if "-" in word and re.match(r"^(Di|Tri|Tetra|Penta|Hexa|Bis|Tris)", word):
        return True
    # N,N- or N- prefixed systematic name (N,N-Dibutyl-formamide)
    if re.match(r"^N[,N]*-[A-Z]", word) and "-" in word[3:]:
        return True
    # Contains [ ] ring notation (Benzo[1,3]dioxole)
    if "[" in word:
        return True
    # Lowercase positional prefix + hyphen (o-Tolylamine, p-Xylene, m-Cresol)
    if re.match(r"^[a-z]-[A-Z]", word):
        return True
    return False


def _looks_like_iupac_name(word):
    """Check if a single unhyphenated word is likely a systematic IUPAC name."""
    # Must have a recognized stem AND suffix
    has_stem = bool(_IUPAC_STEMS.search(word))
    has_suffix = bool(_IUPAC_SUFFIX.search(word))
    if has_stem and has_suffix:
        return True
    # Known heterocyclic / systematic names without standard stem+suffix
    known = {
        "aziridine", "thiirane", "oxirane", "thietane", "oxetane",
        "anthraquinone", "benzoquinone", "quinone",
    }
    if word.lower() in known:
        return True
    return False


def _has_iupac_suffix(word):
    """Check if a word ends with a recognized chemical suffix (lenient check)."""
    return bool(_IUPAC_SUFFIX.search(word))


def extract_common_name(raw_name):
    """Extract the common (trivial) name from 'CommonName IUPACName' format in Hansen A1."""
    name = raw_name.strip()
    if not name:
        return name

    # 1. Handle asterisk delimiter (explicit marker in the source)
    if "*" in name:
        return name.split("*")[0].strip()

    # 2. Strip parenthetical aliases — but only when the part before the paren
    #    is a meaningful name (at least 3 chars), to avoid stripping systematic
    #    notation like "2-(Diethylamino)".
    paren_match = re.match(r"^(.+?)\s*\([^)]*\)\s*(.+)$", name)
    if paren_match:
        before = paren_match.group(1).strip()
        after = paren_match.group(2).strip()
        if after and len(before) >= 3 and not before.endswith("-"):
            return before

    words = name.split()
    n = len(words)

    if n <= 1:
        return name

    # 3. Check for exact word-for-word duplication (case-insensitive)
    # "Acetamide Acetamide", "Benzoic Acid Benzoic acid"
    for half_len in range(1, n // 2 + 1):
        if n >= 2 * half_len:
            first_half = " ".join(words[:half_len])
            second_half = " ".join(words[half_len : 2 * half_len])
            if first_half.lower() == second_half.lower():
                return first_half

    # 4. Find IUPAC boundary by pattern matching on hyphenated/digit-starting words
    for i in range(1, n):
        word = words[i]

        # IUPAC starts with digit (3-Chloro-propene, 2-Methoxy-phenylamine)
        if word[0].isdigit():
            return " ".join(words[:i])

        # IUPAC starts with parenthetical ((E)-But-2-enal, (Z)-1-Bromo-propene)
        if word[0] in "([":
            return " ".join(words[:i])

        # IUPAC word has systematic hyphenation
        if _looks_like_iupac_word(word):
            return " ".join(words[:i])

        # Lowercase word after capitalized context → IUPAC continuation
        # e.g., "Butyric Anhydride Butanoic anhydride" → "anhydride" is lowercase
        if word[0].islower() and i >= 2:
            prev = words[i - 1]
            if prev[0].isupper():
                return " ".join(words[:i - 1])

    # 5. For 3-word entries: check if last word is a single IUPAC name
    # e.g., "Dimethyl Ether Methoxymethane", "Methyl Bromide Bromomethane"
    if n == 3:
        if _looks_like_iupac_name(words[2]):
            return " ".join(words[:2])
        # Also catch compound words with just a suffix (Allylamine, Vinylamine)
        if _has_iupac_suffix(words[2]) and words[2][0].isupper():
            return " ".join(words[:2])

    # 6. For 2-word entries: check if second word is a plausible IUPAC name
    if n == 2:
        if _looks_like_iupac_name(words[1]):
            return words[0]
        # Keep both words for compound names (Castor Oil, Pine Oil)
        return name

    # 7. Fallback: return the full name
    return name


def normalize_cas(cas):
    """Clean up CAS number."""
    if not cas:
        return ""
    cas = str(cas).strip()
    if re.match(r"^\d{2,7}-\d{2}-\d$", cas):
        return cas
    return ""


def parse_float(val):
    """Parse a float value, returning None for missing/invalid data."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if val != val:  # NaN check
            return None
        return float(val)
    val = str(val).strip()
    if not val or val == "-" or val.lower() in ("nan", "none", ""):
        return None
    try:
        return float(val.replace(",", ""))
    except ValueError:
        return None


def classify_chemical(name, smiles=None):
    """Simple rule-based classification based on name/SMILES patterns."""
    name_lower = name.lower() if name else ""

    # Alcohols (check before hydrocarbons due to overlap like "cyclohexanol")
    if any(x in name_lower for x in [
        "methanol", "ethanol", "propanol", "butanol", "pentanol", "hexanol",
        "octanol", "heptanol", "nonanol", "decanol", "alcohol", "glycol",
        "glycerol", "phenol", "cresol", "naphthol", "catechol", "resorcinol",
        "hydroquinone", "cyclohexanol", "menthol", "borneol", "fenchol",
        "isoborneol", "terpineol",
    ]):
        return "alcohol"

    # Amides (check before amines)
    if any(x in name_lower for x in [
        "formamide", "acetamide", "pyrrolidone", "pyrrolidinone",
        "nmp", "dmf", "dmac", "caprolactam", "dimethylformamide",
        "dimethylacetamide",
    ]):
        return "amide"

    # Glycol ethers (check before ethers)
    if any(x in name_lower for x in [
        "methoxyethanol", "ethoxyethanol", "butoxyethanol", "propoxyethanol",
        "glycol ether", "cellosolve", "carbitol", "dowanol",
        "methoxypropanol", "ethoxypropanol",
        "diethylene glycol mono", "propylene glycol mono",
    ]):
        return "glycol ether"

    # Esters (check before acids)
    if any(x in name_lower for x in [
        "acetate", "formate", "propionate", "butyrate", "benzoate",
        "lactone", "butyrolactone", "propiolactone", "valerolactone",
        "carbonate", "acrylate", "methacrylate", "lactate",
        "malonate", "oxalate", "succinate", "phthalate", "maleate",
        "fumarate", "citrate", "tartrate", "sebacate", "adipate",
    ]):
        return "ester"

    # Ketones
    if any(x in name_lower for x in [
        "acetone", "ketone", "cyclohexanone", "cyclopentanone",
        "acetophenone", "isophorone", "mesityl oxide", "diacetone",
        "methylethylketone", "methyl ethyl ketone", "mek", "mibk",
        "methyl isobutyl ketone",
    ]):
        return "ketone"

    # Ethers
    if any(x in name_lower for x in [
        "ether", "tetrahydrofuran", "thf", "dioxane", "dioxolane",
        "diglyme", "triglyme", "tetraglyme", "glyme", "furan",
        "anisole", "phenetole", "methyltetrahydrofuran",
        "dihydropyran", "tetrahydropyran",
    ]):
        return "ether"

    # Nitriles
    if any(x in name_lower for x in [
        "nitrile", "cyanide", "acetonitrile", "propionitrile",
        "butyronitrile", "benzonitrile", "acrylonitrile", "succinonitrile",
    ]):
        return "nitrile"

    # Amines
    if any(x in name_lower for x in [
        "amine", "aniline", "pyridine", "morpholine", "triethylamine",
        "ethanolamine", "piperidine", "piperazine", "imidazole",
        "diethylamine", "dimethylamine", "trimethylamine", "butylamine",
        "hexylamine", "diisopropylamine", "dibutylamine",
        "cyclohexylamine", "benzylamine",
    ]):
        return "amine"

    # Sulfoxides / sulfones
    if any(x in name_lower for x in [
        "sulfoxide", "dmso", "sulfolane", "sulfone",
        "dimethyl sulfoxide", "dimethylsulfoxide",
    ]):
        return "sulfoxide"

    # Acids
    if any(x in name_lower for x in [
        "acetic acid", "formic acid", "propionic acid", "butyric acid",
        "valeric acid", "caproic acid", "oleic acid", "stearic acid",
        "benzoic acid", "lactic acid", "citric acid", "oxalic acid",
        "trifluoroacetic", " acid",
    ]):
        return "acid"

    # Nitro compounds
    if any(x in name_lower for x in [
        "nitromethane", "nitroethane", "nitropropane", "nitrobenzene",
        "nitrotoluene", "dinitro", "trinitro",
    ]):
        if name_lower.startswith("nitro") or "nitro" in name_lower:
            return "nitro"

    # Terpenes
    if any(x in name_lower for x in [
        "limonene", "pinene", "cymene", "terpene", "turpentine",
        "myrcene", "camphene", "carvone", "geraniol", "linalool",
        "eucalyptol", "camphor", "thymol", "cineole",
    ]):
        return "terpene"

    # Halogenated
    if any(x in name_lower for x in [
        "chlor", "brom", "iodo", "fluoro",
    ]):
        # Exclude fluorinated compounds (separate category)
        if any(x in name_lower for x in ["perfluoro", "hexafluoro", "trifluoroethanol"]):
            return "fluorinated"
        return "halogenated"

    # Fluorinated
    if any(x in name_lower for x in [
        "perfluoro", "hexafluoro", "trifluoro", "fluorinated",
    ]):
        return "fluorinated"

    # Aromatics
    if any(x in name_lower for x in [
        "benzene", "toluene", "xylene", "styrene", "naphthalene",
        "tetralin", "ethylbenzene", "phenyl", "biphenyl", "anthracene",
        "cumene", "mesitylene", "indene", "fluorene", "acenaphthene",
        "azulene", "durene", "hemimellitene",
    ]):
        return "aromatic"

    # Hydrocarbons
    if any(x in name_lower for x in [
        "hexane", "heptane", "octane", "pentane", "decane", "nonane",
        "undecane", "dodecane", "butane", "propane",
        "cyclohexane", "cyclopentane", "methylcyclohexane",
        "naphtha", "paraffin", "isooctane", "petroleum", "squalane",
    ]):
        return "hydrocarbon"

    # Heterocyclic
    if any(x in name_lower for x in [
        "furfural", "furfuryl", "pyrrole", "thiophene", "oxazole",
        "thiazole", "indole", "quinoline", "isoquinoline", "carbazole",
        "acridine",
    ]):
        return "heterocyclic"

    # Glycols (without ether)
    if "glycol" in name_lower and "ether" not in name_lower:
        return "glycol"

    # Inorganic
    if name_lower in ["water", "carbon disulfide", "carbon disulphide"]:
        return "inorganic"

    # Aldehyde
    if any(x in name_lower for x in [
        "aldehyde", "formaldehyde", "acetaldehyde", "propionaldehyde",
        "butyraldehyde", "benzaldehyde", "furfural",
    ]):
        return "aldehyde"

    # Thiol / sulfide
    if any(x in name_lower for x in [
        "mercaptan", "thiol", "sulfide", "disulfide", "thioacet",
    ]):
        return "sulfur compound"

    # Fall back to SMILES-based classification
    if smiles:
        s = smiles
        has_ring = "c" in s or "C1" in s
        has_O = "O" in s
        has_N = "N" in s
        has_halogen = any(x in s for x in ["Cl", "Br", "I", "F"])

        if has_halogen and not has_O:
            return "halogenated"
        if has_ring and not has_O and not has_N and not has_halogen:
            return "aromatic"
        if not has_O and not has_N and not has_halogen and not has_ring:
            return "hydrocarbon"

    return "other"


# ---------------------------------------------------------------------------
# Data loading functions
# ---------------------------------------------------------------------------


def load_original_seed():
    """Load the original curated seed data."""
    sys.path.insert(0, BASE_DIR)
    from backend.app.data.seed_data import SOLVENTS, POLYMERS

    chemicals = []
    for row in SOLVENTS:
        name, cas, dd, dp, dh, mw, bp, density, mv, cat = row
        chemicals.append({
            "name": name,
            "cas_number": cas,
            "smiles": "",
            "molecular_formula": "",
            "delta_d": dd, "delta_p": dp, "delta_h": dh,
            "molecular_weight": mw,
            "boiling_point": bp,
            "density": density,
            "molar_volume": mv,
            "category": cat,
            "ghs_hazard": "",
            "source": "handbook",
            "source_url": SOURCE_URLS["handbook"],
        })

    polymers = []
    for row in POLYMERS:
        name, dd, dp, dh, r0, ptype = row
        polymers.append({
            "name": name,
            "delta_d": dd, "delta_p": dp, "delta_h": dh,
            "radius": r0,
            "type": ptype,
            "cas_number": "",
            "source": "handbook",
            "source_url": SOURCE_URLS["handbook"],
        })

    return chemicals, polymers


def load_mendeley():
    """Load Mendeley dataset (499 solvents)."""
    filepath = os.path.join(RAW_DIR, "mendeley_hsp.csv")
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return []

    chemicals = []
    with open(filepath, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = row.get("Solvent", "").strip()
            if not name:
                continue
            dd = parse_float(row.get("dD_MPa05"))
            dp = parse_float(row.get("dP_MPa05"))
            dh = parse_float(row.get("dH_MPa05"))
            if dd is None or dp is None or dh is None:
                continue

            ghs = row.get("GHS", "").strip()
            h_stmts = row.get("H_Statements", "").strip()
            ghs_info = ""
            if h_stmts:
                ghs_info = h_stmts

            chemicals.append({
                "name": name,
                "cas_number": normalize_cas(row.get("CAS", "")),
                "smiles": row.get("SMILES", "").strip(),
                "molecular_formula": "",
                "delta_d": dd, "delta_p": dp, "delta_h": dh,
                "molecular_weight": parse_float(row.get("MWt_g_mol")),
                "boiling_point": parse_float(row.get("Tb_C")),
                "density": parse_float(row.get("Density_g_cm3")),
                "molar_volume": parse_float(row.get("MVol_cm3_mol")),
                "category": "",
                "ghs_hazard": ghs_info,
                "source": "mendeley",
                "source_url": SOURCE_URLS["mendeley"],
            })

    return chemicals


def load_solvpred():
    """Load SolvPred dataset (249 solvents)."""
    filepath = os.path.join(RAW_DIR, "solvpred_hsp.json")
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return []

    with open(filepath, "r") as f:
        data = json.load(f)

    chemicals = []
    for entry in data:
        name = entry.get("Name", "").strip()
        if not name:
            continue
        dd = parse_float(entry.get("D"))
        dp = parse_float(entry.get("P"))
        dh = parse_float(entry.get("H"))
        if dd is None or dp is None or dh is None:
            continue

        chemicals.append({
            "name": name,
            "cas_number": normalize_cas(str(entry.get("CAS", ""))),
            "smiles": str(entry.get("SMILES", "") or "").strip(),
            "molecular_formula": "",
            "delta_d": dd, "delta_p": dp, "delta_h": dh,
            "molecular_weight": parse_float(entry.get("mw")),
            "boiling_point": parse_float(entry.get("bp")),
            "density": None,
            "molar_volume": parse_float(entry.get("Mole_vol")),
            "category": "",
            "ghs_hazard": "",
            "source": "solvpred",
            "source_url": SOURCE_URLS["solvpred"],
        })

    return chemicals


def load_accudyne_solvents():
    """Load Accudyne solvents (89 entries)."""
    filepath = os.path.join(RAW_DIR, "accudyne_solvents.csv")
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return []

    chemicals = []
    with open(filepath, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = row.get("Name", "").strip()
            if not name:
                continue
            dd = parse_float(row.get("\u03b4D (MPa^0.5)"))
            dp = parse_float(row.get("\u03b4P (MPa^0.5)"))
            dh = parse_float(row.get("\u03b4H (MPa^0.5)"))
            if dd is None or dp is None or dh is None:
                continue

            chemicals.append({
                "name": name,
                "cas_number": normalize_cas(row.get("CAS Number", "")),
                "smiles": "",
                "molecular_formula": row.get("Molecular Formula", "").strip(),
                "delta_d": dd, "delta_p": dp, "delta_h": dh,
                "molecular_weight": parse_float(row.get("Molecular Weight")),
                "boiling_point": None,
                "density": None,
                "molar_volume": parse_float(row.get("Molar Volume (cm\u00b3/mol)")),
                "category": "",
                "ghs_hazard": "",
                "source": "accudyne",
                "source_url": SOURCE_URLS["accudyne"],
            })

    return chemicals


def load_accudyne_polymers():
    """Load Accudyne polymers (those with HSP data)."""
    filepath = os.path.join(RAW_DIR, "accudyne_polymers.csv")
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return []

    polymers = []
    with open(filepath, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = row.get("Name", "").strip()
            if not name:
                continue
            dd = parse_float(row.get("\u03b4D (MPa^0.5)"))
            dp = parse_float(row.get("\u03b4P (MPa^0.5)"))
            dh = parse_float(row.get("\u03b4H (MPa^0.5)"))
            if dd is None or dp is None or dh is None:
                continue

            polymers.append({
                "name": name,
                "delta_d": dd, "delta_p": dp, "delta_h": dh,
                "radius": parse_float(row.get("R0 (MPa^0.5)")),
                "type": "",
                "cas_number": normalize_cas(row.get("CAS Number", "")),
                "source": "accudyne",
                "source_url": "https://www.accudynetest.com/polytable_02.html",
            })

    return polymers


def load_wolfram():
    """Load Wolfram dataset (211 solvents)."""
    filepath = os.path.join(RAW_DIR, "wolfram_hsp.csv")
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return []

    chemicals = []
    with open(filepath, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = row.get("Solvent", "").strip()
            if not name:
                continue
            dd = parse_float(row.get("delta_d_MPa0.5"))
            dp = parse_float(row.get("delta_p_MPa0.5"))
            dh = parse_float(row.get("delta_h_MPa0.5"))
            if dd is None or dp is None or dh is None:
                continue

            chemicals.append({
                "name": name,
                "cas_number": "",
                "smiles": "",
                "molecular_formula": "",
                "delta_d": dd, "delta_p": dp, "delta_h": dh,
                "molecular_weight": None,
                "boiling_point": None,
                "density": None,
                "molar_volume": parse_float(row.get("Volume_cm3_per_mol")),
                "category": "",
                "ghs_hazard": "",
                "source": "wolfram",
                "source_url": SOURCE_URLS["wolfram"],
            })

    return chemicals


def load_hansen_1k():
    """Load Hansen 1k dataset (1,183 compounds)."""
    filepath = os.path.join(RAW_DIR, "hansen_1k_smiles_shorter.csv")
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return []

    chemicals = []
    with open(filepath, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = row.get("Molecule", "").strip()
            if not name:
                continue
            dd = parse_float(row.get("hansen_d"))
            dp = parse_float(row.get("hansen_p"))
            dh = parse_float(row.get("hansen_h"))
            if dd is None or dp is None or dh is None:
                continue

            # Some entries have CAS numbers as names
            cas = ""
            if re.match(r"^\d{2,7}-\d{2}-\d$", name):
                cas = name
                name = f"CAS {name}"

            chemicals.append({
                "name": name,
                "cas_number": cas,
                "smiles": row.get("SMILES", "").strip(),
                "molecular_formula": "",
                "delta_d": dd, "delta_p": dp, "delta_h": dh,
                "molecular_weight": None,
                "boiling_point": None,
                "density": None,
                "molar_volume": None,
                "category": "",
                "ghs_hazard": "",
                "source": "pang2024",
                "source_url": SOURCE_URLS["pang2024"],
            })

    return chemicals


def load_hspip_solvents():
    """Load solvents from the HSPiP software database (HSPiPsolvents.xlsx).

    Columns: Chemical name, CAS #, Dispersion, Polarity, H-H bonding
    """
    try:
        import openpyxl
    except ImportError:
        print("  Warning: openpyxl not installed, skipping HSPiP data")
        return []

    filepath = os.path.join(RAW_DIR, "HSPiPsolvents.xlsx")
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return []

    chemicals = []
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))  # skip header
    wb.close()

    for row in rows:
        if not row or len(row) < 5:
            continue
        name = str(row[0]).strip() if row[0] else ""
        cas = str(row[1]).strip() if row[1] else ""
        if not name:
            continue

        dd = parse_float(row[2])
        dp = parse_float(row[3])
        dh = parse_float(row[4])
        if dd is None or dp is None or dh is None:
            continue

        chemicals.append({
            "name": name,
            "cas_number": cas,
            "smiles": "",
            "molecular_formula": "",
            "delta_d": dd, "delta_p": dp, "delta_h": dh,
            "molecular_weight": None,
            "boiling_point": None,
            "density": None,
            "molar_volume": None,
            "category": "",
            "ghs_hazard": "",
            "source": "hspip",
            "source_url": SOURCE_URLS["hspip"],
        })

    return chemicals


def load_hansen_a2():
    """Load Hansen Appendix Table A.2 polymers from cleaned OCR output.

    New CSV columns: number, polymer_name, category, dispersion, polar,
    hydrogen_bonding, interaction_radius
    """
    filepath = os.path.join(OUT_DIR, "table_a2.csv")
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return []

    polymers = []
    with open(filepath, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = row.get("polymer_name", "").strip()
            if not name:
                continue

            dd = parse_float(row.get("dispersion"))
            dp = parse_float(row.get("polar"))
            dh = parse_float(row.get("hydrogen_bonding"))
            if dd is None or dp is None or dh is None:
                continue

            polymers.append({
                "name": name,
                "delta_d": dd, "delta_p": dp, "delta_h": dh,
                "radius": parse_float(row.get("interaction_radius")),
                "type": row.get("category", "").strip(),
                "cas_number": "",
                "source": "hansen_a2",
                "source_url": SOURCE_URLS["hansen_a2"],
            })

    return polymers


# ---------------------------------------------------------------------------
# Merging
# ---------------------------------------------------------------------------


def _merge_metadata(existing, new):
    """Fill in missing metadata from new source without overwriting HSP values."""
    for field in ["cas_number", "smiles", "molecular_formula", "molecular_weight",
                  "boiling_point", "density", "molar_volume", "ghs_hazard"]:
        if not existing.get(field) and new.get(field):
            existing[field] = new[field]
    # Track how many independent sources corroborate this entry
    if "source_count" not in existing:
        existing["source_count"] = 1
    new_src = new.get("source", "")
    if new_src and new_src != existing.get("source", ""):
        existing["source_count"] = existing.get("source_count", 1) + 1


def merge_chemicals(all_sources):
    """Merge chemical datasets, deduplicating by CAS then by normalized name."""
    by_cas = {}
    by_name = {}
    result = []

    for chem in all_sources:
        cas = chem["cas_number"]
        norm = normalize_name(chem["name"])

        # Check for duplicate by CAS
        if cas and cas in by_cas:
            _merge_metadata(by_cas[cas], chem)
            continue

        # Check for duplicate by normalized name
        if norm and norm in by_name:
            _merge_metadata(by_name[norm], chem)
            # Also register CAS if the new source has one
            if cas and cas not in by_cas:
                by_cas[cas] = by_name[norm]
            continue

        # New compound
        result.append(chem)
        if cas:
            by_cas[cas] = chem
        if norm:
            by_name[norm] = chem

    return result


def merge_polymers(all_sources):
    """Merge polymer datasets, deduplicating by normalized name."""
    by_name = {}
    result = []

    for poly in all_sources:
        norm = normalize_name(poly["name"])

        if norm and norm in by_name:
            existing = by_name[norm]
            if not existing.get("cas_number") and poly.get("cas_number"):
                existing["cas_number"] = poly["cas_number"]
            if not existing.get("radius") and poly.get("radius"):
                existing["radius"] = poly["radius"]
            if not existing.get("type") and poly.get("type"):
                existing["type"] = poly["type"]
            continue

        result.append(poly)
        if norm:
            by_name[norm] = poly

    return result


def classify_polymer(name):
    """Classify polymer type based on name, using specific type labels
    matching the section headers from the Hansen Appendix A.2 table."""
    nl = name.lower()

    # Specific polymer types (ordered from most to least specific)
    if any(x in nl for x in ["epoxy", "epoxies", "epon", "epikote"]):
        return "Epoxy"
    if any(x in nl for x in ["polyurethane", "polyurethanes"]):
        return "Polyurethane"
    if any(x in nl for x in ["polyimide", "polyamideimide"]):
        return "Polyimide"
    if any(x in nl for x in ["nylon", "polyamide", "pa6", "pa11", "pa12"]):
        return "Polyamide"
    if any(x in nl for x in ["polycarbonate"]):
        return "Polycarbonate"
    if any(x in nl for x in ["polysulfone", "polysulphone"]):
        return "Polysulfone PSU"
    if any(x in nl for x in ["polystyrene", "abs"]):
        return "Polystyrene"
    if any(x in nl for x in ["polyvinyl chloride", "pvc"]):
        return "Polyvinylchloride"
    if any(x in nl for x in ["polyvinyl acetate", "pvac"]):
        return "Polyvinylacetate"
    if any(x in nl for x in ["polyvinyl alcohol", "pva)", "pvoh"]):
        return "Polyvinyl Alcohol"
    if any(x in nl for x in ["polyvinyl butyral", "pvb"]):
        return "Polyvinylbutyral"
    if any(x in nl for x in ["polyvinylidene", "pvdc", "saran"]):
        return "Polyvinylidene Chloride"
    if any(x in nl for x in ["polyvinylpyrrolidone", "pvp"]):
        return "Polyvinylpyrrolidone"
    if any(x in nl for x in ["pmma", "pema", "pibma", "pbma", "methacrylate",
                               "acrylate", "acrylic", "plexiglas"]):
        return "Polyacrylate"
    if any(x in nl for x in ["polyacrylonitrile", "pan)"]):
        return "Polyacrylonitrile"
    if any(x in nl for x in ["polyethylene terephthalate", "pet)", "pla)",
                               "polylactic"]):
        return "Polyester"
    if any(x in nl for x in ["polyethylene", "hdpe", "ldpe", "pe)"]):
        return "Polyethylene"
    if any(x in nl for x in ["polypropylene", "pp)"]):
        return "Polypropylene"
    if any(x in nl for x in ["polyisoprene", "pip)"]):
        return "Polyisoprene"
    if any(x in nl for x in ["polybutadiene"]):
        return "Polybutadiene"
    if any(x in nl for x in ["polyphenylene oxide", "ppo"]):
        return "Polyphenylene Oxide"
    if any(x in nl for x in ["polyphenylene sulfide", "pps"]):
        return "Polyphenylene Sulfide"
    if any(x in nl for x in ["polyetherimide", "pei)"]):
        return "Polyetherimide"
    if any(x in nl for x in ["polyethersulfone", "pes)"]):
        return "Polyethersulfone"
    if any(x in nl for x in ["polychlorotrifluoroethylene", "pctfe",
                               "ptfe", "pvdf", "fluoropolymer", "teflon",
                               "fep", "fluorinated ethylene", "pfa"]):
        return "Fluoropolymer"
    if any(x in nl for x in ["phenolic"]):
        return "Phenolic Resins"
    if any(x in nl for x in ["alkyd"]):
        return "Alkyd"
    if any(x in nl for x in ["rubber", "elastomer", "neoprene", "epdm"]):
        return "Elastomer"
    if any(x in nl for x in ["silicone", "pdms"]):
        return "Silicone Resins"
    if any(x in nl for x in ["cellulose", "cellophane"]):
        return "Cellulose"
    if any(x in nl for x in ["starch", "lignin", "shellac", "rosin",
                               "bitumen", "natural", "chitosan", "chitin",
                               "zein", "collagen"]):
        return "Natural"
    if any(x in nl for x in ["nitrocellulose"]):
        return "Nitrocellulose"
    return ""


# ---------------------------------------------------------------------------
# Confidence scoring
# ---------------------------------------------------------------------------

# Source reliability tiers (higher = more reliable)
SOURCE_CONFIDENCE = {
    "handbook": 0.50,    # Curated seed data from Hansen Handbook
    "hspip": 0.50,       # HSPiP software database (authoritative, with CAS)
    "mendeley": 0.40,    # Peer-reviewed dataset (Langner & Brabec 2022)
    "solvpred": 0.35,    # Academic tool (Fang et al.)
    "accudyne": 0.40,    # Manufacturer/test lab data
    "wolfram": 0.35,     # Curated data repository
    "pang2024": 0.30,    # HSPiP database extract
    "hansen_a2": 0.30,   # OCR from Hansen Appendix A.2
}


def compute_confidence(entry, is_polymer=False):
    """Compute a confidence score (0.0-1.0) for a database entry.

    Scoring factors:
      - Base: source reliability tier (0.30-0.50)
      - CAS number present: +0.15 (verified chemical identity)
      - SMILES present: +0.10 (structural confirmation, chemicals only)
      - Cross-referenced (>1 source): +0.15 per additional source (max +0.30)
    """
    source = entry.get("source", "")
    score = SOURCE_CONFIDENCE.get(source, 0.25)

    # CAS number = verified identity
    if entry.get("cas_number"):
        score += 0.15

    # SMILES = structural confirmation (chemicals only)
    if not is_polymer and entry.get("smiles"):
        score += 0.10

    # Cross-referencing bonus: each additional source adds confidence
    source_count = entry.get("source_count", 1)
    if source_count > 1:
        score += min((source_count - 1) * 0.15, 0.30)

    return round(min(score, 1.0), 2)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main():
    os.makedirs(OUT_DIR, exist_ok=True)

    print("Building comprehensive HSP database...")
    print()
    print("Loading data sources:")

    original_chems, original_polys = load_original_seed()
    print(f"  Original seed:    {len(original_chems):>5} chemicals, {len(original_polys)} polymers")

    mendeley_chems = load_mendeley()
    print(f"  Mendeley:         {len(mendeley_chems):>5} chemicals")

    solvpred_chems = load_solvpred()
    print(f"  SolvPred:         {len(solvpred_chems):>5} chemicals")

    accudyne_chems = load_accudyne_solvents()
    print(f"  Accudyne solv:    {len(accudyne_chems):>5} chemicals")

    accudyne_polys = load_accudyne_polymers()
    print(f"  Accudyne poly:    {len(accudyne_polys):>5} polymers")

    wolfram_chems = load_wolfram()
    print(f"  Wolfram:          {len(wolfram_chems):>5} chemicals")

    hansen_1k_chems = load_hansen_1k()
    print(f"  Hansen 1k:        {len(hansen_1k_chems):>5} chemicals")

    hspip_chems = load_hspip_solvents()
    print(f"  HSPiP solvents:   {len(hspip_chems):>5} chemicals")

    hansen_a2_polys = load_hansen_a2()
    print(f"  Hansen A2:        {len(hansen_a2_polys):>5} polymers")

    total_raw = (len(original_chems) + len(mendeley_chems) + len(solvpred_chems)
                 + len(accudyne_chems) + len(wolfram_chems) + len(hansen_1k_chems)
                 + len(hspip_chems))
    print(f"  ---")
    print(f"  Total raw:        {total_raw:>5} entries")

    # Merge in priority order
    print()
    print("Merging and deduplicating...")
    all_chems = (original_chems + mendeley_chems + solvpred_chems
                 + accudyne_chems + wolfram_chems + hansen_1k_chems
                 + hspip_chems)
    merged_chems = merge_chemicals(all_chems)

    all_polys = original_polys + accudyne_polys + hansen_a2_polys
    merged_polys = merge_polymers(all_polys)

    # Classify unclassified chemicals
    for chem in merged_chems:
        if not chem["category"]:
            chem["category"] = classify_chemical(chem["name"], chem.get("smiles"))

    # Classify untyped polymers
    for poly in merged_polys:
        if not poly.get("type"):
            poly["type"] = classify_polymer(poly["name"])

    print(f"  Result: {len(merged_chems)} unique chemicals, {len(merged_polys)} unique polymers")

    # Enrich CAS numbers from cache (built by scripts/enrich_cas.py)
    cas_cache_path = os.path.join(OUT_DIR, ".cas_cache.json")
    if os.path.exists(cas_cache_path):
        import json as _json
        with open(cas_cache_path) as _f:
            cas_cache = _json.load(_f)
        enriched = 0
        for chem in merged_chems:
            if not chem.get("cas_number"):
                cached = cas_cache.get(chem["name"].lower().strip(), "")
                if cached:
                    chem["cas_number"] = cached
                    enriched += 1
        if enriched:
            print(f"  CAS enrichment from cache: {enriched} additional CAS numbers")

    # Compute confidence scores (after CAS enrichment so CAS bonus is accurate)
    for chem in merged_chems:
        chem["confidence"] = compute_confidence(chem, is_polymer=False)
    for poly in merged_polys:
        poly["confidence"] = compute_confidence(poly, is_polymer=True)

    # Sort
    merged_chems.sort(key=lambda x: x["name"].lower())
    merged_polys.sort(key=lambda x: x["name"].lower())

    # Write chemicals CSV
    chem_path = os.path.join(OUT_DIR, "hsp_chemicals.csv")
    chem_fields = [
        "name", "cas_number", "smiles", "molecular_formula",
        "delta_d", "delta_p", "delta_h",
        "molecular_weight", "boiling_point", "density", "molar_volume",
        "category", "ghs_hazard", "confidence", "source_count", "source", "source_url",
    ]

    with open(chem_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=chem_fields, extrasaction="ignore")
        writer.writeheader()
        for chem in merged_chems:
            row = {}
            for field in chem_fields:
                val = chem.get(field)
                row[field] = val if val is not None else ""
            writer.writerow(row)

    print(f"  Written: {chem_path}")

    # Write polymers CSV
    poly_path = os.path.join(OUT_DIR, "hsp_polymers.csv")
    poly_fields = [
        "name", "cas_number", "delta_d", "delta_p", "delta_h",
        "radius", "type", "confidence", "source_count", "source", "source_url",
    ]

    with open(poly_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=poly_fields, extrasaction="ignore")
        writer.writeheader()
        for poly in merged_polys:
            row = {}
            for field in poly_fields:
                val = poly.get(field)
                row[field] = val if val is not None else ""
            writer.writerow(row)

    print(f"  Written: {poly_path}")

    # Summary
    print()
    print("=" * 50)
    print(f"TOTAL CHEMICALS: {len(merged_chems)}")
    print(f"TOTAL POLYMERS:  {len(merged_polys)}")
    print("=" * 50)

    # Source breakdown
    source_counts = {}
    for chem in merged_chems:
        src = chem.get("source", "unknown")
        source_counts[src] = source_counts.get(src, 0) + 1
    print("\nChemicals by primary source:")
    for src, count in sorted(source_counts.items(), key=lambda x: -x[1]):
        url = SOURCE_URLS.get(src, "")
        print(f"  {src:12s}: {count:>5}  ({url})")

    # Category breakdown
    cat_counts = {}
    for chem in merged_chems:
        cat = chem.get("category", "other")
        cat_counts[cat] = cat_counts.get(cat, 0) + 1
    print("\nChemicals by category:")
    for cat, count in sorted(cat_counts.items(), key=lambda x: -x[1]):
        print(f"  {cat:20s}: {count:>5}")

    # Metadata completeness
    has_cas = sum(1 for c in merged_chems if c.get("cas_number"))
    has_smiles = sum(1 for c in merged_chems if c.get("smiles"))
    has_mw = sum(1 for c in merged_chems if c.get("molecular_weight"))
    has_bp = sum(1 for c in merged_chems if c.get("boiling_point"))
    n = len(merged_chems)
    print(f"\nMetadata completeness:")
    print(f"  CAS number:   {has_cas:>5}/{n} ({100*has_cas/n:.0f}%)")
    print(f"  SMILES:       {has_smiles:>5}/{n} ({100*has_smiles/n:.0f}%)")
    print(f"  Mol. weight:  {has_mw:>5}/{n} ({100*has_mw/n:.0f}%)")
    print(f"  Boiling point:{has_bp:>5}/{n} ({100*has_bp/n:.0f}%)")


if __name__ == "__main__":
    main()
